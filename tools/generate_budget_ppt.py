"""Generate charts, analysis suggestions, and a PPT deck for budget-vs-actual review.

Usage:
  python tools/generate_budget_ppt.py \
    --input data_sample_budget_actual.csv \
    --output-dir output/2026-03
"""

from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.util import Inches


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="CSV path with required columns")
    parser.add_argument("--output-dir", required=True, help="Output folder")
    return parser.parse_args()


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    rename_map = {}
    if "actual" not in df.columns and "actual_amount" in df.columns:
        rename_map["actual_amount"] = "actual"
    if "budget" not in df.columns and "budget_amount" in df.columns:
        rename_map["budget_amount"] = "budget"

    if rename_map:
        df = df.rename(columns=rename_map)

    required_cols = {"period", "module", "actual", "budget"}
    missing = required_cols - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")

    return df


def build_monthly_chart(df: pd.DataFrame, chart_path: Path) -> None:
    monthly = (
        df.groupby("period", as_index=False)[["actual", "budget"]]
        .sum()
        .sort_values("period")
    )

    plt.figure(figsize=(10, 5))
    x = monthly["period"].astype(str)
    plt.plot(x, monthly["actual"], marker="o", label="Actual")
    plt.plot(x, monthly["budget"], marker="o", label="Budget")
    plt.title("Budget vs Actual (Monthly)")
    plt.xlabel("Period")
    plt.ylabel("Amount")
    plt.legend()
    plt.tight_layout()
    plt.savefig(chart_path, dpi=150)
    plt.close()


def build_module_chart(df: pd.DataFrame, chart_path: Path) -> None:
    module_df = df.groupby("module", as_index=False)[["actual", "budget"]].sum()

    plt.figure(figsize=(10, 5))
    x = range(len(module_df))
    width = 0.35
    plt.bar([i - width / 2 for i in x], module_df["budget"], width=width, label="Budget")
    plt.bar([i + width / 2 for i in x], module_df["actual"], width=width, label="Actual")
    plt.xticks(list(x), module_df["module"], rotation=30, ha="right")
    plt.title("Budget vs Actual (By Module)")
    plt.ylabel("Amount")
    plt.legend()
    plt.tight_layout()
    plt.savefig(chart_path, dpi=150)
    plt.close()


def generate_suggestions(df: pd.DataFrame) -> list[str]:
    group = df.groupby("module", as_index=False)[["actual", "budget"]].sum()
    group["var_amt"] = group["actual"] - group["budget"]
    group["var_rate"] = group["var_amt"] / group["budget"].abs().replace(0, pd.NA)

    suggestions: list[str] = []
    for row in group.sort_values("var_amt", ascending=False).itertuples(index=False):
        module = str(row.module)
        var_amt = float(row.var_amt)
        var_rate = row.var_rate

        if pd.isna(var_rate):
            suggestions.append(f"{module}: 预算为 0，建议复核预算口径与拆分逻辑。")
            continue

        rate_pct = float(var_rate) * 100
        if rate_pct >= 10:
            suggestions.append(
                f"{module}: 实际高于预算 {rate_pct:.1f}%（+{var_amt:,.0f}），建议做价格/销量/结构三因素归因，并设置费用上限。"
            )
        elif rate_pct <= -10:
            suggestions.append(
                f"{module}: 实际低于预算 {abs(rate_pct):.1f}%（{var_amt:,.0f}），建议检查收入确认节奏、交付进度或预算目标合理性。"
            )
        else:
            suggestions.append(
                f"{module}: 偏差 {rate_pct:.1f}%（{var_amt:,.0f}），总体可控，建议持续跟踪滚动预测。"
            )

    if not suggestions:
        suggestions.append("暂无可分析数据。")

    return suggestions


def save_suggestions(suggestions: list[str], output_path: Path) -> None:
    content = "\n".join(f"- {item}" for item in suggestions)
    output_path.write_text(content, encoding="utf-8")


def build_ppt(
    monthly_chart_path: Path,
    module_chart_path: Path,
    suggestions: list[str],
    out_ppt: Path,
) -> None:
    prs = Presentation()

    cover_layout = prs.slide_layouts[5]
    cover = prs.slides.add_slide(cover_layout)
    cover_title = cover.shapes.title
    if cover_title is not None:
        cover_title.text = "Budget Review"

    cover.shapes.add_picture(str(monthly_chart_path), Inches(0.8), Inches(1.3), width=Inches(8.0))

    module_slide = prs.slides.add_slide(cover_layout)
    module_title = module_slide.shapes.title
    if module_title is not None:
        module_title.text = "Module Analysis"
    module_slide.shapes.add_picture(str(module_chart_path), Inches(0.8), Inches(1.3), width=Inches(8.0))

    suggest_slide = prs.slides.add_slide(cover_layout)
    suggest_title = suggest_slide.shapes.title
    if suggest_title is not None:
        suggest_title.text = "Analysis Suggestions"

    text_box = suggest_slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.5), Inches(4.8))
    tf = text_box.text_frame
    tf.clear()
    for index, item in enumerate(suggestions):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = f"• {item}"

    prs.save(out_ppt)


def main() -> None:
    args = parse_args()
    output_dir = Path(args.output_dir)
    chart_dir = output_dir / "charts"
    chart_dir.mkdir(parents=True, exist_ok=True)

    df = pd.read_csv(args.input)
    df = normalize_columns(df)

    monthly_chart_path = chart_dir / "budget_vs_actual_monthly.png"
    module_chart_path = chart_dir / "budget_vs_actual_module.png"
    suggestion_path = output_dir / "analysis_suggestions.md"
    ppt_path = output_dir / "Budget_Review.pptx"

    build_monthly_chart(df, monthly_chart_path)
    build_module_chart(df, module_chart_path)
    suggestions = generate_suggestions(df)
    save_suggestions(suggestions, suggestion_path)
    build_ppt(monthly_chart_path, module_chart_path, suggestions, ppt_path)

    print(f"Chart generated: {monthly_chart_path}")
    print(f"Chart generated: {module_chart_path}")
    print(f"Suggestions generated: {suggestion_path}")
    print(f"PPT generated: {ppt_path}")


if __name__ == "__main__":
    main()
