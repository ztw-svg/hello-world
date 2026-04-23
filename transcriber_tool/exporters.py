from __future__ import annotations

import importlib.util
from pathlib import Path

from transcriber_tool.models import TranscriptSegment
from transcriber_tool.utils import format_timestamp


def _role_tag(seg: TranscriptSegment) -> str:
    tags: list[str] = []
    if seg.speaker:
        tags.append(seg.speaker)
    if seg.gender:
        tags.append(seg.gender)
    if not tags:
        return ""
    return "[" + "][".join(tags) + "] "


def render_lines(segments: list[TranscriptSegment], timestamp_format: str, include_timestamps: bool) -> list[str]:
    lines: list[str] = []
    for seg in segments:
        role = _role_tag(seg)
        body = f"{role}{seg.text}".strip()
        if include_timestamps:
            lines.append(f"{format_timestamp(seg.start, timestamp_format)} {body}")
        else:
            lines.append(body)
    return lines


def export_markdown(path: Path, title: str, lines: list[str]) -> None:
    content = [f"# {title}", "", *lines, ""]
    path.write_text("\n".join(content), encoding="utf-8")


def export_txt(path: Path, title: str, lines: list[str]) -> None:
    content = [title, "=" * len(title), "", *lines, ""]
    path.write_text("\n".join(content), encoding="utf-8")


def export_docx(path: Path, title: str, lines: list[str]) -> None:
    if importlib.util.find_spec("docx") is None:
        raise RuntimeError("导出 DOCX 需要 python-docx，请安装 requirements.txt")
    from docx import Document

    doc = Document()
    doc.add_heading(title, 1)
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
